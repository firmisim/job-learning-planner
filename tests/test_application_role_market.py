from __future__ import annotations

import inspect
import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from application.errors import ApplicationError
from application.market import MarketOperations
from schemas.core import AtomicMarketSignal, JDAnalysesDocument, JDAnalysis
from src.market import jd_source_fingerprint


def _xlsx(rows: list[tuple[str, str, str | None]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Jobs"
    sheet.append(["job_title", "jd_text", "company"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _analysis_document(operations: MarketOperations, role_id: str, signals: list[AtomicMarketSignal]) -> str:
    job = operations.storage.load_jds(operations._uuid(role_id, "role_id")).jobs[0]
    return JDAnalysesDocument(
        role_id=job.role_id,
        analyses=[JDAnalysis(job_id=job.job_id, role_id=job.role_id, source_fingerprint=jd_source_fingerprint(job), signals=signals)],
    ).model_dump_json()


def test_simple_role_lifecycle_preserves_uuid_and_isolation(tmp_path: Path) -> None:
    operations = MarketOperations(tmp_path / "state")
    first_token = operations.view().roles_sha256
    role_a, _ = operations.create_role("AI 应用开发", first_token)
    role_b, _ = operations.create_role("Java 后端", operations.view(role_a).roles_sha256)
    original_id = role_a

    operations.rename_role(role_a, "AI 产品开发", operations.view(role_a).roles_sha256)
    assert operations.view(role_a).current_role is not None
    assert operations.view(role_a).current_role.role_id == original_id
    assert operations.view(role_a).current_role.name == "AI 产品开发"

    global_capability = operations.storage.create_capability("Python")
    operations.add_job(role_a, title="AI Engineer", company=None, jd_text="使用 FastAPI", source_url=None, expected_sha256=operations.view(role_a).jds_sha256 or "")
    operations.add_job(role_b, title="Java Engineer", company=None, jd_text="使用 Spring Boot", source_url=None, expected_sha256=operations.view(role_b).jds_sha256 or "")
    assert [job.title for job in operations.view(role_a).jobs] == ["AI Engineer"]
    assert [job.title for job in operations.view(role_b).jobs] == ["Java Engineer"]

    operations.delete_role(role_a, confirmation="AI 产品开发", expected_sha256=operations.view(role_a).roles_sha256)
    assert operations.storage.load_catalog().capabilities == [global_capability]
    assert [role.name for role in operations.storage.load_roles().roles] == ["Java 后端"]


def test_jd_add_import_replace_delete_are_role_scoped(tmp_path: Path) -> None:
    operations = MarketOperations(tmp_path / "state")
    role_a, _ = operations.create_role("A", operations.view().roles_sha256)
    role_b, _ = operations.create_role("B", operations.view(role_a).roles_sha256)
    operations.add_job(role_a, title="Old", company=None, jd_text="Old text", source_url=None, expected_sha256=operations.view(role_a).jds_sha256 or "")
    operations.add_job(role_b, title="Other", company=None, jd_text="Other text", source_url=None, expected_sha256=operations.view(role_b).jds_sha256 or "")

    operations.import_jobs(role_a, "jobs.xlsx", _xlsx([("Added", "New text", None)]), mode="append", expected_sha256=operations.view(role_a).jds_sha256 or "")
    assert [item.title for item in operations.view(role_a).jobs] == ["Old", "Added"]
    operations.import_jobs(role_a, "jobs.xlsx", _xlsx([("Replacement", "Spring Boot 或 FastAPI", "Example")]), mode="replace", expected_sha256=operations.view(role_a).jds_sha256 or "")
    assert [item.title for item in operations.view(role_a).jobs] == ["Replacement"]
    assert [item.title for item in operations.view(role_b).jobs] == ["Other"]

    job_id = operations.view(role_a).jobs[0].job_id
    foreign_job_id = operations.view(role_b).jobs[0].job_id
    role_a_sha256 = operations.view(role_a).jds_sha256 or ""
    with pytest.raises(ApplicationError) as foreign_selection:
        operations.delete_jobs(
            role_a,
            [foreign_job_id],
            expected_sha256=role_a_sha256,
        )
    assert foreign_selection.value.category == "validation"
    assert [item.title for item in operations.view(role_a).jobs] == ["Replacement"]

    with pytest.raises(ApplicationError) as stale_selection:
        operations.delete_jobs(role_a, [job_id], expected_sha256="0" * 64)
    assert stale_selection.value.category == "conflict"
    assert [item.title for item in operations.view(role_a).jobs] == ["Replacement"]

    operations.delete_jobs(role_a, [job_id], expected_sha256=operations.view(role_a).jds_sha256 or "")
    assert operations.view(role_a).jobs == ()
    assert len(operations.view(role_b).jobs) == 1


def test_analysis_handoff_normalizes_copy_without_changing_raw_jd(tmp_path: Path) -> None:
    operations = MarketOperations(tmp_path / "state")
    role_id, _ = operations.create_role("Backend", operations.view().roles_sha256)
    raw_text = "  Apply now\r\n\r\n\r\n职责：\t使用  FastAPI\x00\r福利：年度体检  "
    operations.add_job(
        role_id,
        title="Backend",
        company=None,
        jd_text=raw_text,
        source_url=None,
        expected_sha256=operations.view(role_id).jds_sha256 or "",
    )
    stored_before = operations.storage.load_jds(
        operations._uuid(role_id, "role_id")
    ).jobs[0].jd_text

    operations.prepare_analysis(role_id)
    request = operations.handoffs.load_request("jd-analysis")
    prepared_job = request["input"]["jobs"][0]

    stored_job = operations.storage.load_jds(
        operations._uuid(role_id, "role_id")
    ).jobs[0]
    assert stored_job.jd_text == stored_before
    assert "\r" in stored_job.jd_text
    assert "\x00" in stored_job.jd_text
    assert prepared_job["jd_text"] == (
        "Apply now\n\n职责： 使用 FastAPI\n福利：年度体检"
    )
    assert "Apply now" in prepared_job["jd_text"]


def test_analysis_lineage_atomic_cases_and_jd_deduped_statistics(tmp_path: Path) -> None:
    operations = MarketOperations(tmp_path / "state")
    role_id, _ = operations.create_role("Backend", operations.view().roles_sha256)
    text = "熟悉 Spring Boot 或 FastAPI；掌握 Python、Java或者C/C++；理解 CI/CD、TCP/IP 与 I/O。"
    operations.add_job(role_id, title="Backend", company=None, jd_text=text, source_url=None, expected_sha256=operations.view(role_id).jds_sha256 or "")
    signals = [
        AtomicMarketSignal(source_expression="Spring Boot 或 FastAPI", atomic_expression="Spring Boot", evidence="熟悉 Spring Boot 或 FastAPI"),
        AtomicMarketSignal(source_expression="Spring Boot 或 FastAPI", atomic_expression="FastAPI", evidence="熟悉 Spring Boot 或 FastAPI"),
        *(AtomicMarketSignal(source_expression="Python、Java或者C/C++", atomic_expression=name, evidence="掌握 Python、Java或者C/C++") for name in ("Python", "Java", "C/C++")),
    ]
    signals.extend([
        AtomicMarketSignal(source_expression="CI/CD", atomic_expression="CI/CD", evidence="理解 CI/CD、TCP/IP 与 I/O"),
        AtomicMarketSignal(source_expression="TCP/IP", atomic_expression="TCP/IP", evidence="理解 CI/CD、TCP/IP 与 I/O"),
        AtomicMarketSignal(source_expression="I/O", atomic_expression="I/O", evidence="理解 CI/CD、TCP/IP 与 I/O"),
        AtomicMarketSignal(source_expression="FastAPI", atomic_expression="FastAPI", evidence="Spring Boot 或 FastAPI"),
    ])
    payload = _analysis_document(operations, role_id, signals)
    operations.import_analyses(role_id, payload, expected_sha256=operations.view(role_id).analyses_sha256 or "")

    view = operations.view(role_id)
    assert {item.atomic_expression for item in view.signals} == {"Spring Boot", "FastAPI", "Python", "Java", "C/C++", "CI/CD", "TCP/IP", "I/O"}
    fastapi = next(item for item in view.signals if item.atomic_expression == "FastAPI")
    assert fastapi.jd_count == 1
    assert fastapi.sample_size == 1
    assert {item.source_expression for item in fastapi.evidence} == {"Spring Boot 或 FastAPI", "FastAPI"}
    operations.add_job(role_id, title="Pending", company=None, jd_text="尚未分析", source_url=None, expected_sha256=view.jds_sha256 or "")
    partial = operations.view(role_id)
    assert partial.jd_count == 2
    assert partial.sample_size == 1
    assert next(item for item in partial.signals if item.atomic_expression == "FastAPI").sample_size == 1


def test_malformed_or_stale_analysis_is_rejected_and_reanalysis_replaces(tmp_path: Path) -> None:
    operations = MarketOperations(tmp_path / "state")
    role_id, _ = operations.create_role("Backend", operations.view().roles_sha256)
    operations.add_job(role_id, title="Backend", company=None, jd_text="使用 FastAPI", source_url=None, expected_sha256=operations.view(role_id).jds_sha256 or "")
    first = _analysis_document(operations, role_id, [AtomicMarketSignal(source_expression="FastAPI", atomic_expression="FastAPI", evidence="使用 FastAPI")])
    operations.import_analyses(role_id, first, expected_sha256=operations.view(role_id).analyses_sha256 or "")
    second = _analysis_document(operations, role_id, [AtomicMarketSignal(source_expression="FastAPI", atomic_expression="API development", evidence="使用 FastAPI")])
    operations.import_analyses(role_id, second, expected_sha256=operations.view(role_id).analyses_sha256 or "")
    assert [item.atomic_expression for item in operations.view(role_id).signals] == ["API development"]

    malformed = json.loads(second)
    malformed["analyses"][0]["signals"][0]["evidence"] = "not in JD"
    with pytest.raises(ApplicationError) as error:
        operations.import_analyses(role_id, json.dumps(malformed), expected_sha256=operations.view(role_id).analyses_sha256 or "")
    assert error.value.category == "validation"


def test_full_reanalysis_can_atomically_replace_an_unreadable_generated_asset(
    tmp_path: Path,
) -> None:
    operations = MarketOperations(tmp_path / "state")
    role_id, _ = operations.create_role("Backend", operations.view().roles_sha256)
    operations.add_job(
        role_id,
        title="API",
        company=None,
        jd_text="使用 FastAPI",
        source_url=None,
        expected_sha256=operations.view(role_id).jds_sha256 or "",
    )
    operations.add_job(
        role_id,
        title="Runtime",
        company=None,
        jd_text="使用 Python",
        source_url=None,
        expected_sha256=operations.view(role_id).jds_sha256 or "",
    )
    selected = operations._uuid(role_id, "role_id")
    jobs = operations.storage.load_jds(selected).jobs
    analysis_path = tmp_path / "state" / "roles" / role_id / "jd_analyses.json"
    incompatible = {
        "role_id": role_id,
        "analyses": [
            {
                "job_id": str(jobs[0].job_id),
                "role_id": role_id,
                "source_fingerprint": jd_source_fingerprint(jobs[0]),
                "signals": [
                    {
                        "source_expression": "FastAPI",
                        "atomic_expression": "FastAPI",
                        "evidence": "使用 FastAPI",
                        "requirement_type": "required",
                    }
                ],
            }
        ],
    }
    analysis_path.write_text(json.dumps(incompatible), encoding="utf-8")
    original_bytes = analysis_path.read_bytes()

    operations.prepare_analysis(role_id)
    request = operations.handoffs.load_request("jd-analysis")
    assert request["context"]["expected_sha256"] == operations.storage.jd_analyses_sha256(selected)

    partial = JDAnalysesDocument(
        role_id=selected,
        analyses=[
            JDAnalysis(
                job_id=jobs[0].job_id,
                role_id=selected,
                source_fingerprint=jd_source_fingerprint(jobs[0]),
                signals=[
                    AtomicMarketSignal(
                        source_expression="FastAPI",
                        atomic_expression="FastAPI",
                        evidence="使用 FastAPI",
                    )
                ],
            )
        ],
    )
    with pytest.raises(ApplicationError) as error:
        operations.import_analyses(
            role_id,
            partial.model_dump_json(),
            expected_sha256=str(request["context"]["expected_sha256"]),
        )
    assert error.value.category == "validation"
    assert analysis_path.read_bytes() == original_bytes

    complete = JDAnalysesDocument(
        role_id=selected,
        analyses=[
            JDAnalysis(
                job_id=job.job_id,
                role_id=selected,
                source_fingerprint=jd_source_fingerprint(job),
                signals=[
                    AtomicMarketSignal(
                        source_expression=expression,
                        atomic_expression=expression,
                        evidence=f"使用 {expression}",
                    )
                ],
            )
            for job, expression in zip(jobs, ("FastAPI", "Python"), strict=True)
        ],
    )
    operations.import_analyses(
        role_id,
        complete.model_dump_json(),
        expected_sha256=str(request["context"]["expected_sha256"]),
    )

    assert {item.atomic_expression for item in operations.view(role_id).signals} == {
        "FastAPI",
        "Python",
    }
    assert b"requirement_type" not in analysis_path.read_bytes()


def test_stage_b_market_path_has_no_taxonomy_governance_or_mechanical_split() -> None:
    import application.market as application_market
    import src.market as market_core

    source = inspect.getsource(application_market) + inspect.getsource(market_core)
    assert "governance" not in source.casefold()
    assert "taxonomy" not in source.casefold()
    assert "re.split" not in source
    assert "Direction" not in source
