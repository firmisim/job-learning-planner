from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from io import BytesIO
from typing import Iterable
from uuid import UUID

from openpyxl import load_workbook

from schemas.core import JD, JDAnalysis


def jd_source_fingerprint(jd: JD) -> str:
    """Fingerprint only the source fields that semantic analysis consumes."""

    payload = {
        "company": jd.company,
        "jd_text": jd.jd_text,
        "source_url": str(jd.source_url) if jd.source_url else None,
        "title": jd.title,
    }
    encoded = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def jd_duplicate_key(jd: JD) -> str:
    return jd_source_fingerprint(jd)


def analysis_is_current(jd: JD, analysis: JDAnalysis | None) -> bool:
    return bool(
        analysis
        and analysis.job_id == jd.job_id
        and analysis.role_id == jd.role_id
        and analysis.source_fingerprint == jd_source_fingerprint(jd)
    )


def validate_analysis_against_jd(jd: JD, analysis: JDAnalysis) -> None:
    if analysis.job_id != jd.job_id or analysis.role_id != jd.role_id:
        raise ValueError("JDAnalysis provenance 与 JD 不一致")
    if analysis.source_fingerprint != jd_source_fingerprint(jd):
        raise ValueError("JDAnalysis source_fingerprint 已过期")
    normalized_source = normalize_jd_analysis_text(jd.jd_text)
    for signal in analysis.signals:
        if normalize_jd_analysis_text(signal.evidence) not in normalized_source:
            raise ValueError("JDAnalysis evidence 必须是 JD 原文中的连续片段")


def normalize_atomic_expression(value: str) -> str:
    return " ".join(value.split()).casefold()


def normalize_jd_analysis_text(value: str) -> str:
    """Remove only low-risk formatting artifacts from an analysis-time copy."""

    normalized = value.replace("\r\n", "\n").replace("\r", "\n")
    normalized = "".join(
        character
        for character in normalized
        if character in {"\n", "\t"} or character.isprintable()
    )
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in normalized.split("\n")]
    compacted: list[str] = []
    for line in lines:
        if line or not compacted or compacted[-1]:
            compacted.append(line)
    return "\n".join(compacted).strip()


@dataclass(frozen=True)
class SignalEvidence:
    job_id: UUID
    job_title: str
    source_expression: str
    evidence: str


@dataclass(frozen=True)
class AtomicSignalStatistic:
    atomic_expression: str
    jd_count: int
    sample_size: int
    evidence: tuple[SignalEvidence, ...]


def build_atomic_market_statistics(
    jobs: Iterable[JD], analyses: Iterable[JDAnalysis]
) -> tuple[AtomicSignalStatistic, ...]:
    job_list = list(jobs)
    analysis_by_id = {item.job_id: item for item in analyses}
    grouped: dict[str, dict[str, object]] = {}
    for job in job_list:
        analysis = analysis_by_id.get(job.job_id)
        if not analysis_is_current(job, analysis):
            continue
        assert analysis is not None
        seen_in_job: set[str] = set()
        for signal in analysis.signals:
            key = normalize_atomic_expression(signal.atomic_expression)
            group = grouped.setdefault(
                key,
                {
                    "label": signal.atomic_expression.strip(),
                    "job_ids": set(),
                    "evidence": [],
                },
            )
            evidence_key = (
                str(job.job_id),
                signal.source_expression,
                signal.evidence,
            )
            evidence_items = group["evidence"]
            assert isinstance(evidence_items, list)
            if evidence_key not in {
                (str(item.job_id), item.source_expression, item.evidence)
                for item in evidence_items
            }:
                evidence_items.append(
                    SignalEvidence(
                        job_id=job.job_id,
                        job_title=job.title,
                        source_expression=signal.source_expression,
                        evidence=signal.evidence,
                    )
                )
            if key not in seen_in_job:
                job_ids = group["job_ids"]
                assert isinstance(job_ids, set)
                job_ids.add(job.job_id)
                seen_in_job.add(key)
    sample_size = sum(
        1
        for job in job_list
        if analysis_is_current(job, analysis_by_id.get(job.job_id))
    )
    result = []
    for value in grouped.values():
        job_ids = value["job_ids"]
        evidence = value["evidence"]
        assert isinstance(job_ids, set) and isinstance(evidence, list)
        result.append(
            AtomicSignalStatistic(
                atomic_expression=str(value["label"]),
                jd_count=len(job_ids),
                sample_size=sample_size,
                evidence=tuple(evidence),
            )
        )
    return tuple(
        sorted(
            result,
            key=lambda item: (-item.jd_count, item.atomic_expression.casefold()),
        )
    )


_HEADER_ALIASES = {
    "job_title": "title",
    "title": "title",
    "company": "company",
    "jd_text": "jd_text",
    "source_url": "source_url",
}


def read_jds_workbook(content: bytes, role_id: UUID) -> list[JD]:
    """Read the one supported structured JD input without a generic import layer."""

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 .xlsx JD workbook") from exc
    sheet = workbook["Jobs"] if "Jobs" in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise ValueError("JD workbook 为空") from exc
    headers: dict[int, str] = {}
    for index, value in enumerate(raw_headers):
        key = str(value or "").strip().casefold()
        if key in _HEADER_ALIASES:
            headers[index] = _HEADER_ALIASES[key]
    if "title" not in headers.values() or "jd_text" not in headers.values():
        raise ValueError("JD workbook 必须包含 job_title/title 与 jd_text 列")
    jobs: list[JD] = []
    for row_number, row in enumerate(rows, start=2):
        values = {
            field: str(row[index]).strip()
            for index, field in headers.items()
            if index < len(row) and row[index] is not None and str(row[index]).strip()
        }
        if not values:
            continue
        try:
            jobs.append(
                JD(
                    role_id=role_id,
                    title=values.get("title", ""),
                    company=values.get("company"),
                    jd_text=values.get("jd_text", ""),
                    source_url=values.get("source_url"),
                )
            )
        except Exception as exc:
            raise ValueError(f"JD workbook 第 {row_number} 行无效: {exc}") from exc
    if not jobs:
        raise ValueError("JD workbook 没有可导入的 JD")
    keys = [jd_duplicate_key(item) for item in jobs]
    if len(keys) != len(set(keys)):
        raise ValueError("JD workbook 含重复 JD")
    return jobs


def analysis_input_document(role_id: UUID, jobs: Iterable[JD]) -> str:
    payload = {
        "schema_version": "1.0",
        "role_id": str(role_id),
        "jobs": [
            {
                **item.model_dump(mode="json"),
                "jd_text": normalize_jd_analysis_text(item.jd_text),
                "source_fingerprint": jd_source_fingerprint(item),
            }
            for item in jobs
        ],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True)
